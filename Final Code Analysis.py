import tkinter as tk
from tkinter import ttk, filedialog, messagebox, colorchooser
import pandas as pd
import openpyxl
from openpyxl.styles import PatternFill
import re
import copy

class AnalysisApp(tk.Tk):
    AUTO_YELLOW = "#FFFF00"  # Yellow highlight color for modified cells

    def __init__(self):
        super().__init__()
        self.title("Analysis With Original to Formatted — Enterprise Edition")
        self.geometry("1420x880")
        self.minsize(1150, 720)

        # Application State
        self.file_path = None
        self.workbook_sheets = []
        self.orig_sheet_name = None
        self.fmt_sheet_name = None

        self.df_orig = pd.DataFrame()
        self.filtered_df_orig = pd.DataFrame()

        self.df_fmt = pd.DataFrame()
        self.filtered_df_fmt = pd.DataFrame()

        self.selected_row_idx = None
        self.selected_col_name = None
        self.active_selection_type = None  # 'row', 'col', or 'cell'

        self.history = []          # Stack of dicts for undo
        self.cell_highlights = {}  # Maps (row_idx, col_name) -> hex_color

        self._apply_styles()
        self._build_ui()

    def _apply_styles(self):
        style = ttk.Style(self)
        style.theme_use("clam")
        style.configure("Treeview.Heading", font=("Helvetica", 9, "bold"), background="#e1e1e1")
        style.configure("Treeview", rowheight=24, font=("Helvetica", 9))
        style.map("Treeview", background=[("selected", "#0078d7")], foreground=[("selected", "white")])

    # =========================================================================
    # 1. MAIN APPLICATION LAYOUT
    # =========================================================================
    def _build_ui(self):
        # Top Action Bar
        top_frame = ttk.Frame(self, padding=10)
        top_frame.pack(fill=tk.X)

        self.btn_load = ttk.Button(top_frame, text="1. Select Excel Workbook", command=self.load_workbook)
        self.btn_load.pack(side=tk.LEFT, padx=5)

        self.lbl_file = ttk.Label(top_frame, text="No workbook loaded", font=("Helvetica", 9, "italic"))
        self.lbl_file.pack(side=tk.LEFT, padx=10)

        self.btn_sheets = ttk.Button(top_frame, text="2. Select Excel Sheets", command=self.open_sheet_select_popup, state=tk.DISABLED)
        self.btn_sheets.pack(side=tk.LEFT, padx=5)

        self.btn_export = ttk.Button(top_frame, text="Get Final Export File", command=self.export_final_file, state=tk.DISABLED)
        self.btn_export.pack(side=tk.RIGHT, padx=5)

        # Task Control Panel Dropdowns
        task_frame = ttk.LabelFrame(self, text="Task Operations Menu", padding=10)
        task_frame.pack(fill=tk.X, padx=10, pady=5)

        # 1. Task By Row
        ttk.Label(task_frame, text="Select Task By Row:").grid(row=0, column=0, padx=5, sticky=tk.W)
        self.row_task_var = tk.StringVar(value="Choose Task...")
        self.row_task_cb = ttk.Combobox(
            task_frame, textvariable=self.row_task_var, state="readonly", width=22,
            values=["Choose Task...", "Highlight Selected Row With Color", "Add Selected Duplicated Row", "Edit Entire Row"]
        )
        self.row_task_cb.grid(row=0, column=1, padx=5)
        self.row_task_cb.bind("<<ComboboxSelected>>", self.handle_row_task)

        # 2. Task By Cell / Column
        ttk.Label(task_frame, text="Select Task By Cell/Column:").grid(row=0, column=2, padx=5, sticky=tk.W)
        self.col_task_var = tk.StringVar(value="Choose Task...")
        self.col_task_cb = ttk.Combobox(
            task_frame, textvariable=self.col_task_var, state="readonly", width=22,
            values=["Choose Task...", "Highlight Selected Cell With Color", "Edit Cell", "Delimited Cell Text"]
        )
        self.col_task_cb.grid(row=0, column=3, padx=5)
        self.col_task_cb.bind("<<ComboboxSelected>>", self.handle_col_task)

        # 3. Other Task (Includes Analysis Full Data Pop-Up Trigger)
        ttk.Label(task_frame, text="Select Other Task:").grid(row=0, column=4, padx=5, sticky=tk.W)
        self.other_task_var = tk.StringVar(value="Choose Task...")
        self.other_task_cb = ttk.Combobox(
            task_frame, textvariable=self.other_task_var, state="readonly", width=24,
            values=[
                "Choose Task...", "Analysis Full Data", "Find Data", "Find & Replace", "Insert Blank Row", 
                "Copy Cell Value", "Trim Extra Spaces", "Change Text Case", "CountIf Summary",
                "Batch: Keep First Pipe Value", "Batch: Keep Second Pipe Value", 
                "Fill Blank Cells", "Remove Duplicate Rows",
                "Undo Last Task", "View Task History"
            ]
        )
        self.other_task_cb.grid(row=0, column=5, padx=5)
        self.other_task_cb.bind("<<ComboboxSelected>>", self.handle_other_task)

        # Dynamic Context Metadata Display Bar
        meta_frame = ttk.Frame(self, padding=(10, 5))
        meta_frame.pack(fill=tk.X)
        self.lbl_metadata = ttk.Label(
            meta_frame, 
            text="Selection Metadata: Click any Row, Column Header, or Cell in the Formatted Data sheet to inspect.",
            font=("Consolas", 9, "bold"), foreground="#003366", wraplength=1380
        )
        self.lbl_metadata.pack(side=tk.LEFT)

        # Workspace Dual View Splitter
        paned = ttk.PanedWindow(self, orient=tk.HORIZONTAL)
        paned.pack(fill=tk.BOTH, expand=True, padx=10, pady=5)

        # Original Data Sheet (Read Only View with Filtering)
        orig_container = ttk.LabelFrame(paned, text="Original / Raw Data Sheet (View Only Reference)")
        paned.add(orig_container, weight=1)

        filter_bar_orig = ttk.Frame(orig_container, padding=2)
        filter_bar_orig.pack(fill=tk.X)
        ttk.Label(filter_bar_orig, text="Filter Column:").pack(side=tk.LEFT, padx=2)
        self.filter_col_orig_cb = ttk.Combobox(filter_bar_orig, state="readonly", width=14)
        self.filter_col_orig_cb.pack(side=tk.LEFT, padx=2)
        ttk.Label(filter_bar_orig, text="Search:").pack(side=tk.LEFT, padx=2)
        self.filter_ent_orig = ttk.Entry(filter_bar_orig, width=12)
        self.filter_ent_orig.pack(side=tk.LEFT, padx=2)
        ttk.Button(filter_bar_orig, text="Apply Filter", command=self.apply_orig_filter).pack(side=tk.LEFT, padx=2)
        ttk.Button(filter_bar_orig, text="Clear", command=self.clear_orig_filter).pack(side=tk.LEFT, padx=2)

        self.tree_orig = self._build_treeview(orig_container, is_formatted=False)

        # Formatted Data Sheet (Interactive Working Copy Table)
        fmt_container = ttk.LabelFrame(paned, text="Formatted Data Sheet (Interactive Working Copy)")
        paned.add(fmt_container, weight=1)

        filter_bar_fmt = ttk.Frame(fmt_container, padding=2)
        filter_bar_fmt.pack(fill=tk.X)
        ttk.Label(filter_bar_fmt, text="Filter Column:").pack(side=tk.LEFT, padx=2)
        self.filter_col_fmt_cb = ttk.Combobox(filter_bar_fmt, state="readonly", width=14)
        self.filter_col_fmt_cb.pack(side=tk.LEFT, padx=2)
        ttk.Label(filter_bar_fmt, text="Search:").pack(side=tk.LEFT, padx=2)
        self.filter_ent_fmt = ttk.Entry(filter_bar_fmt, width=12)
        self.filter_ent_fmt.pack(side=tk.LEFT, padx=2)
        ttk.Button(filter_bar_fmt, text="Apply Filter", command=self.apply_formatted_filter).pack(side=tk.LEFT, padx=2)
        ttk.Button(filter_bar_fmt, text="Clear", command=self.clear_formatted_filter).pack(side=tk.LEFT, padx=2)

        self.tree_fmt = self._build_treeview(fmt_container, is_formatted=True)

        # History Audit Trail Bar
        history_frame = ttk.LabelFrame(self, text="Recorded Task Audit History Log", padding=5)
        history_frame.pack(fill=tk.X, padx=10, pady=5)

        self.history_listbox = tk.Listbox(history_frame, height=3, font=("Consolas", 9))
        self.history_listbox.pack(fill=tk.X, side=tk.LEFT, expand=True)

        hist_scroll = ttk.Scrollbar(history_frame, orient=tk.VERTICAL, command=self.history_listbox.yview)
        hist_scroll.pack(side=tk.RIGHT, fill=tk.Y)
        self.history_listbox.config(yscrollcommand=hist_scroll.set)

    def _build_treeview(self, parent, is_formatted=False):
        container = ttk.Frame(parent)
        container.pack(fill=tk.BOTH, expand=True)

        tree = ttk.Treeview(container, show="headings", selectmode="browse")
        vsb = ttk.Scrollbar(container, orient="vertical", command=tree.yview)
        hsb = ttk.Scrollbar(container, orient="horizontal", command=tree.xview)
        tree.configure(yscrollcommand=vsb.set, xscrollcommand=hsb.set)

        tree.grid(row=0, column=0, sticky="nsew")
        vsb.grid(row=0, column=1, sticky="ns")
        hsb.grid(row=1, column=0, sticky="ew")

        container.grid_rowconfigure(0, weight=1)
        container.grid_columnconfigure(0, weight=1)

        if is_formatted:
            tree.bind("<<TreeviewSelect>>", self._on_tree_select)
            tree.bind("<ButtonRelease-1>", self._on_tree_click)

        return tree

    # =========================================================================
    # 2. FILE LOADING & SHEET CONFIGURATION
    # =========================================================================
    def load_workbook(self):
        file_path = filedialog.askopenfilename(filetypes=[("Excel Workbooks", "*.xlsx *.xls")])
        if not file_path:
            return

        self.file_path = file_path
        self.lbl_file.config(text=file_path.split("/")[-1], font=("Helvetica", 9, "bold"))

        try:
            xl = pd.ExcelFile(self.file_path)
            self.workbook_sheets = xl.sheet_names
            self.btn_sheets.config(state=tk.NORMAL)
            self.open_sheet_select_popup()
        except Exception as e:
            messagebox.showerror("File Error", f"Unable to read Excel workbook: {e}")

    def open_sheet_select_popup(self):
        popup = tk.Toplevel(self)
        popup.title("Select Sheet Configuration")
        popup.geometry("420x260")
        popup.grab_set()

        ttk.Label(popup, text="Configure Reference & Formatted Sheets", font=("Helvetica", 11, "bold")).pack(pady=12)

        ttk.Label(popup, text="Original / Raw Data Sheet (View Only):").pack(anchor=tk.W, padx=25)
        orig_var = tk.StringVar(value=self.workbook_sheets[0])
        orig_cb = ttk.Combobox(popup, textvariable=orig_var, values=self.workbook_sheets, state="readonly")
        orig_cb.pack(fill=tk.X, padx=25, pady=4)

        ttk.Label(popup, text="Formatted Data Sheet (Working Copy):").pack(anchor=tk.W, padx=25, pady=(8, 0))
        fmt_default = self.workbook_sheets[1] if len(self.workbook_sheets) > 1 else self.workbook_sheets[0]
        fmt_var = tk.StringVar(value=fmt_default)
        fmt_cb = ttk.Combobox(popup, textvariable=fmt_var, values=self.workbook_sheets, state="readonly")
        fmt_cb.pack(fill=tk.X, padx=25, pady=4)

        def confirm():
            self.orig_sheet_name = orig_var.get()
            self.fmt_sheet_name = fmt_var.get()

            self.df_orig = pd.read_excel(self.file_path, sheet_name=self.orig_sheet_name).fillna("")
            self.filtered_df_orig = self.df_orig.copy()

            self.df_fmt = pd.read_excel(self.file_path, sheet_name=self.fmt_sheet_name).fillna("")
            self.filtered_df_fmt = self.df_fmt.copy()

            self._sync_treeview(self.tree_orig, self.filtered_df_orig)
            self._sync_treeview(self.tree_fmt, self.filtered_df_fmt)

            self.filter_col_orig_cb["values"] = list(self.df_orig.columns)
            if len(self.df_orig.columns) > 0:
                self.filter_col_orig_cb.set(self.df_orig.columns[0])

            self.filter_col_fmt_cb["values"] = list(self.df_fmt.columns)
            if len(self.df_fmt.columns) > 0:
                self.filter_col_fmt_cb.set(self.df_fmt.columns[0])

            self.btn_export.config(state=tk.NORMAL)
            popup.destroy()

        ttk.Button(popup, text="Confirm Sheet Configuration", command=confirm).pack(pady=18)

    def _sync_treeview(self, tree, df):
        tree.delete(*tree.get_children())
        tree["columns"] = list(df.columns)

        for col in df.columns:
            tree.heading(col, text=col, command=lambda c=col: self._on_header_click(c))
            max_len = max([len(str(col))] + [len(str(val)) for val in df[col].head(50)], default=10)
            col_width = max(120, min(max_len * 11, 260))
            tree.column(col, width=col_width, minwidth=110, anchor=tk.W, stretch=False)

        for idx, row in df.iterrows():
            tree.insert("", tk.END, iid=idx, values=list(row))

    def apply_orig_filter(self):
        col = self.filter_col_orig_cb.get()
        query = self.filter_ent_orig.get().strip()
        if not col or not query:
            return

        mask = self.df_orig[col].astype(str).str.contains(query, case=False, regex=False)
        self.filtered_df_orig = self.df_orig[mask].copy()
        self._sync_treeview(self.tree_orig, self.filtered_df_orig)

    def clear_orig_filter(self):
        self.filter_ent_orig.delete(0, tk.END)
        self.filtered_df_orig = self.df_orig.copy()
        self._sync_treeview(self.tree_orig, self.filtered_df_orig)

    def apply_formatted_filter(self):
        col = self.filter_col_fmt_cb.get()
        query = self.filter_ent_fmt.get().strip()
        if not col or not query:
            return

        mask = self.df_fmt[col].astype(str).str.contains(query, case=False, regex=False)
        self.filtered_df_fmt = self.df_fmt[mask].copy()
        self._sync_treeview(self.tree_fmt, self.filtered_df_fmt)

    def clear_formatted_filter(self):
        self.filter_ent_fmt.delete(0, tk.END)
        self.filtered_df_fmt = self.df_fmt.copy()
        self._sync_treeview(self.tree_fmt, self.filtered_df_fmt)

    # =========================================================================
    # 3. INTERACTIVE SELECTION & METADATA COMPUTATION
    # =========================================================================
    def _on_header_click(self, col_name):
        self.selected_col_name = col_name
        self.active_selection_type = 'col'
        self._compute_column_metadata(col_name)

    def _on_tree_click(self, event):
        region = self.tree_fmt.identify_region(event.x, event.y)
        if region == "cell":
            column_id = self.tree_fmt.identify_column(event.x)
            col_index = int(column_id.replace("#", "")) - 1
            selected = self.tree_fmt.selection()
            if selected:
                row_idx = int(selected[0])
                col_name = self.df_fmt.columns[col_index]
                self.selected_row_idx = row_idx
                self.selected_col_name = col_name
                self.active_selection_type = 'cell'
                self._compute_cell_metadata(row_idx, col_name)

    def _on_tree_select(self, event):
        selected = self.tree_fmt.selection()
        if not selected or self.active_selection_type == 'cell':
            return
        row_idx = int(selected[0])
        self.selected_row_idx = row_idx
        self.active_selection_type = 'row'
        self._compute_row_metadata(row_idx)

    def _compute_row_metadata(self, row_idx):
        row_data = self.df_fmt.iloc[row_idx]
        total_headers = len(row_data)
        present_count = sum(1 for v in row_data if str(v).strip() != "")
        empty_count = total_headers - present_count

        pairs = [f"{col}: '{val}'" for col, val in row_data.items() if str(val).strip() != ""]
        values_str = ", ".join(pairs[:4]) + ("..." if len(pairs) > 4 else "")

        meta = f"[ROW {row_idx + 1} METADATA] Headers: {total_headers} | Populated: {present_count} | Empty: {empty_count} | Data: ({values_str})"
        self.lbl_metadata.config(text=meta)

    def _compute_column_metadata(self, col_name):
        col_series = self.df_fmt[col_name].astype(str)
        total_vals = len(col_series)
        empty_count = sum(1 for v in col_series if v.strip() == "")
        value_count = total_vals - empty_count

        spec_char_count = sum(len(re.findall(r'[^a-zA-Z0-9\s]', v)) for v in col_series)
        numeric_count = sum(1 for v in col_series if v.strip().replace('.', '', 1).isdigit())
        dbl_space_count = sum(1 for v in col_series if "  " in v)

        meta = f"[COLUMN '{col_name}' METADATA] Empty Cells: {empty_count} | Populated: {value_count} | Special Chars: {spec_char_count} | Numeric Cells: {numeric_count} | Double Space Rows: {dbl_space_count}"
        self.lbl_metadata.config(text=meta)

    def _compute_cell_metadata(self, row_idx, col_name):
        val = str(self.df_fmt.at[row_idx, col_name])
        length = len(val)
        spec_chars = len(re.findall(r'[^a-zA-Z0-9\s]', val))
        is_num = val.strip().replace('.', '', 1).isdigit()
        has_dbl_space = "Yes" if "  " in val else "No"

        meta = f"[CELL (Row {row_idx + 1}, '{col_name}') METADATA] Value: '{val}' | Length: {length} | Special Chars: {spec_chars} | Is Numeric: {is_num} | Double Spaces: {has_dbl_space}"
        self.lbl_metadata.config(text=meta)

    def _get_target_row(self):
        selected = self.tree_fmt.selection()
        if not selected:
            messagebox.showwarning("Selection Required", "Please click a row in the Formatted Data sheet.")
            return None
        return int(selected[0])

    # =========================================================================
    # 4. AUDIT TRAIL & UNDO ENGINE
    # =========================================================================
    def record_task(self, task_name, details, undo_df=None, undo_hl=None):
        snapshot = {
            "task": task_name,
            "details": details,
            "undo_df": undo_df if undo_df is not None else self.df_fmt.copy(),
            "undo_hl": undo_hl if undo_hl is not None else copy.deepcopy(self.cell_highlights)
        }
        self.history.append(snapshot)
        self.history_listbox.insert(tk.END, f"[{task_name}] {details}")

    def undo_last_task(self):
        if not self.history:
            messagebox.showinfo("Undo Engine", "No recorded operations available to undo.")
            return

        last_state = self.history.pop()
        self.history_listbox.delete(tk.END)

        self.df_fmt = last_state["undo_df"].copy()
        self.cell_highlights = copy.deepcopy(last_state["undo_hl"])
        self.filtered_df_fmt = self.df_fmt.copy()

        self._sync_treeview(self.tree_fmt, self.filtered_df_fmt)
        messagebox.showinfo("Undo Complete", f"Successfully reversed task: {last_state['task']}")

    # =========================================================================
    # 5. TASK BY ROW EXECUTORS
    # =========================================================================
    def handle_row_task(self, event):
        task = self.row_task_var.get()
        if task == "Choose Task...":
            return

        row_idx = self._get_target_row()
        if row_idx is None:
            self.row_task_var.set("Choose Task...")
            return

        if task == "Highlight Selected Row With Color":
            color = colorchooser.askcolor(title="Select Row Color")[1]
            if color:
                undo_df = self.df_fmt.copy()
                undo_hl = copy.deepcopy(self.cell_highlights)
                for col in self.df_fmt.columns:
                    self.cell_highlights[(row_idx, col)] = color
                self.record_task("Highlight Row", f"Applied color {color} to Row {row_idx + 1}", undo_df, undo_hl)

        elif task == "Add Selected Duplicated Row":
            self._popup_edit_row(row_idx, is_duplicate=True)

        elif task == "Edit Entire Row":
            self._popup_edit_row(row_idx, is_duplicate=False)

        self.row_task_var.set("Choose Task...")

    def _popup_edit_row(self, row_idx, is_duplicate=False):
        popup = tk.Toplevel(self)
        title = "Add Duplicated Row" if is_duplicate else f"Edit Entire Row {row_idx + 1}"
        popup.title(title)
        popup.geometry("500x440")
        popup.grab_set()

        canvas = tk.Canvas(popup)
        scrollbar = ttk.Scrollbar(popup, orient="vertical", command=canvas.yview)
        frame = ttk.Frame(canvas)

        frame.bind("<Configure>", lambda e: canvas.configure(scrollregion=canvas.bbox("all")))
        canvas.create_window((0, 0), window=frame, anchor="nw")
        canvas.configure(yscrollcommand=scrollbar.set)

        canvas.pack(side="left", fill="both", expand=True)
        scrollbar.pack(side="right", fill="y")

        entries = {}
        row_data = self.df_fmt.iloc[row_idx]

        ttk.Label(frame, text=title, font=("Helvetica", 10, "bold")).pack(pady=10)

        for col in self.df_fmt.columns:
            f = ttk.Frame(frame)
            f.pack(fill=tk.X, padx=12, pady=3)
            ttk.Label(f, text=str(col), width=20, anchor=tk.W).pack(side=tk.LEFT)
            ent = ttk.Entry(f)
            ent.insert(0, str(row_data[col]))
            ent.pack(side=tk.RIGHT, expand=True, fill=tk.X)
            entries[col] = ent

        def save():
            undo_df = self.df_fmt.copy()
            undo_hl = copy.deepcopy(self.cell_highlights)

            new_row = {col: entries[col].get() for col in self.df_fmt.columns}

            if is_duplicate:
                target_r = len(self.df_fmt)
                self.df_fmt = pd.concat([self.df_fmt, pd.DataFrame([new_row])], ignore_index=True)
                for col in self.df_fmt.columns:
                    self.cell_highlights[(target_r, col)] = self.AUTO_YELLOW
                task_name, desc = "Add Duplicated Row", f"Duplicated Row {row_idx + 1} to index {target_r + 1} (Highlighted Yellow)"
            else:
                for col in self.df_fmt.columns:
                    if str(self.df_fmt.at[row_idx, col]) != new_row[col]:
                        self.df_fmt.at[row_idx, col] = new_row[col]
                        self.cell_highlights[(row_idx, col)] = self.AUTO_YELLOW
                task_name, desc = "Edit Entire Row", f"Updated Row {row_idx + 1} values (Changed cells highlighted Yellow)"

            self.filtered_df_fmt = self.df_fmt.copy()
            self._sync_treeview(self.tree_fmt, self.filtered_df_fmt)
            self.record_task(task_name, desc, undo_df, undo_hl)
            popup.destroy()

        ttk.Button(frame, text="Save Row Data", command=save).pack(pady=15)

    # =========================================================================
    # 6. TASK BY CELL / COLUMN EXECUTORS
    # =========================================================================
    def handle_col_task(self, event):
        task = self.col_task_var.get()
        if task == "Choose Task...":
            return

        row_idx = self._get_target_row()
        if row_idx is None:
            self.col_task_var.set("Choose Task...")
            return

        target_col = self.selected_col_name if self.selected_col_name else self.df_fmt.columns[0]

        if task == "Highlight Selected Cell With Color":
            color = colorchooser.askcolor(title="Select Cell Color")[1]
            if color:
                undo_df = self.df_fmt.copy()
                undo_hl = copy.deepcopy(self.cell_highlights)
                self.cell_highlights[(row_idx, target_col)] = color
                self.record_task("Highlight Cell", f"Highlighted Cell (Row {row_idx + 1}, '{target_col}') with {color}", undo_df, undo_hl)

        elif task == "Edit Cell":
            self._popup_edit_cell(row_idx, target_col)

        elif task == "Delimited Cell Text":
            self._popup_delimited_cell(row_idx, target_col)

        self.col_task_var.set("Choose Task...")

    def _popup_edit_cell(self, row_idx, col_name):
        popup = tk.Toplevel(self)
        popup.title("Edit Cell Value")
        popup.geometry("400x180")
        popup.grab_set()

        curr_val = str(self.df_fmt.at[row_idx, col_name])

        ttk.Label(popup, text=f"Edit Cell (Row {row_idx + 1}, Column '{col_name}')", font=("Helvetica", 10, "bold")).pack(pady=10)
        ent = ttk.Entry(popup, width=42)
        ent.insert(0, curr_val)
        ent.pack(padx=20, pady=5)

        def save():
            undo_df = self.df_fmt.copy()
            undo_hl = copy.deepcopy(self.cell_highlights)
            new_val = ent.get()

            if curr_val != new_val:
                self.df_fmt.at[row_idx, col_name] = new_val
                self.cell_highlights[(row_idx, col_name)] = self.AUTO_YELLOW

            self.filtered_df_fmt = self.df_fmt.copy()
            self._sync_treeview(self.tree_fmt, self.filtered_df_fmt)
            self.record_task("Edit Cell", f"Changed Cell (Row {row_idx + 1}, '{col_name}') -> '{new_val}' (Highlighted Yellow)", undo_df, undo_hl)
            popup.destroy()

        ttk.Button(popup, text="Save Cell Value", command=save).pack(pady=15)

    def _popup_delimited_cell(self, row_idx, col_name):
        popup = tk.Toplevel(self)
        popup.title("Delimited Cell Parsing")
        popup.geometry("450x350")
        popup.grab_set()

        cell_val = str(self.df_fmt.at[row_idx, col_name])

        ttk.Label(popup, text="Delimited Cell Text Processor", font=("Helvetica", 11, "bold")).pack(pady=8)
        ttk.Label(popup, text=f"Source Text: '{cell_val}'").pack(pady=2)

        ttk.Label(popup, text="Enter Delimiter (e.g. , / | .):").pack(anchor=tk.W, padx=25, pady=(10, 0))
        delim_ent = ttk.Entry(popup)
        delim_ent.insert(0, "|")
        delim_ent.pack(fill=tk.X, padx=25, pady=2)

        opts_frame = ttk.Frame(popup)
        opts_frame.pack(fill=tk.BOTH, expand=True, padx=25, pady=10)

        selected_val = tk.StringVar()

        def parse():
            for w in opts_frame.winfo_children():
                w.destroy()
            delim = delim_ent.get()
            if not delim:
                return
            parts = [p.strip() for p in cell_val.split(delim)]
            for idx, part in enumerate(parts):
                rb = ttk.Radiobutton(opts_frame, text=f"Value {idx + 1} — {part}", variable=selected_val, value=part)
                rb.pack(anchor=tk.W, pady=2)
            if parts:
                selected_val.set(parts[0])

        ttk.Button(popup, text="Parse Value", command=parse).pack(pady=2)
        parse()

        def save():
            if not selected_val.get():
                return
            undo_df = self.df_fmt.copy()
            undo_hl = copy.deepcopy(self.cell_highlights)

            chosen = selected_val.get()
            if cell_val != chosen:
                self.df_fmt.at[row_idx, col_name] = chosen
                self.cell_highlights[(row_idx, col_name)] = self.AUTO_YELLOW

            self.filtered_df_fmt = self.df_fmt.copy()
            self._sync_treeview(self.tree_fmt, self.filtered_df_fmt)
            self.record_task("Delimited Cell Text", f"Split Cell (Row {row_idx + 1}, '{col_name}') -> '{chosen}' (Highlighted Yellow)", undo_df, undo_hl)
            popup.destroy()

        ttk.Button(popup, text="Apply Selected Part", command=save).pack(pady=12)

    # =========================================================================
    # 7. OTHER TASK EXECUTORS & ANALYSIS FULL DATA POP-UP
    # =========================================================================
    def handle_other_task(self, event):
        task = self.other_task_var.get()
        if task == "Choose Task...":
            return

        if task == "Analysis Full Data":
            self._open_analysis_full_data_popup()
        elif task == "Find Data":
            self._task_find_data()
        elif task == "Find & Replace":
            self._task_find_replace()
        elif task == "Insert Blank Row":
            self._task_insert_blank_row()
        elif task == "Copy Cell Value":
            self._task_copy_cell_value()
        elif task == "Trim Extra Spaces":
            self._task_trim_spaces()
        elif task == "Change Text Case":
            self._task_change_case()
        elif task == "CountIf Summary":
            self._task_count_if()
        elif task == "Batch: Keep First Pipe Value":
            self._task_batch_pipe_extract(keep_index=0)
        elif task == "Batch: Keep Second Pipe Value":
            self._task_batch_pipe_extract(keep_index=1)
        elif task == "Fill Blank Cells":
            self._task_fill_blanks()
        elif task == "Remove Duplicate Rows":
            self._task_remove_duplicates()
        elif task == "Undo Last Task":
            self.undo_last_task()
        elif task == "View Task History":
            self._task_view_history()

        self.other_task_var.set("Choose Task...")

    def _open_analysis_full_data_popup(self):
        if self.df_fmt.empty:
            messagebox.showwarning("No Data", "Please load and configure an Excel workbook first.")
            return

        popup = tk.Toplevel(self)
        popup.title("Analysis Full Data & Source Inspection")
        popup.geometry("1100x650")
        popup.grab_set()

        # Top Task Menu Frame inside Pop-Up
        task_menu_frame = ttk.LabelFrame(popup, text="Analysis Tasks Menu", padding=10)
        task_menu_frame.pack(fill=tk.X, padx=10, pady=5)

        ttk.Label(task_menu_frame, text="Select Task:").pack(side=tk.LEFT, padx=5)
        analysis_task_var = tk.StringVar(value="Choose Task...")
        analysis_task_cb = ttk.Combobox(
            task_menu_frame, textvariable=analysis_task_var, state="readonly", width=30,
            values=["Choose Task...", "Add Row As is it"]
        )
        analysis_task_cb.pack(side=tk.LEFT, padx=5)

        # Filter Bar for Pop-Up Treeview
        filter_frame = ttk.Frame(popup, padding=5)
        filter_frame.pack(fill=tk.X, padx=10)

        ttk.Label(filter_frame, text="Filter Column:").pack(side=tk.LEFT, padx=2)
        popup_col_cb = ttk.Combobox(filter_frame, values=list(self.df_fmt.columns), state="readonly", width=18)
        popup_col_cb.pack(side=tk.LEFT, padx=2)
        if len(self.df_fmt.columns) > 0:
            popup_col_cb.set(self.df_fmt.columns[0])

        ttk.Label(filter_frame, text="Search:").pack(side=tk.LEFT, padx=2)
        popup_search_ent = ttk.Entry(filter_frame, width=18)
        popup_search_ent.pack(side=tk.LEFT, padx=2)

        display_df = self.df_fmt.copy()

        def apply_popup_filter():
            nonlocal display_df
            col = popup_col_cb.get()
            q = popup_search_ent.get().strip()
            if col and q:
                mask = self.df_fmt[col].astype(str).str.contains(q, case=False, regex=False)
                display_df = self.df_fmt[mask].copy()
            else:
                display_df = self.df_fmt.copy()
            sync_popup_tree(display_df)

        def clear_popup_filter():
            nonlocal display_df
            popup_search_ent.delete(0, tk.END)
            display_df = self.df_fmt.copy()
            sync_popup_tree(display_df)

        ttk.Button(filter_frame, text="Filter", command=apply_popup_filter).pack(side=tk.LEFT, padx=2)
        ttk.Button(filter_frame, text="Clear", command=clear_popup_filter).pack(side=tk.LEFT, padx=2)

        # Treeview Container
        tree_container = ttk.Frame(popup)
        tree_container.pack(fill=tk.BOTH, expand=True, padx=10, pady=5)

        popup_tree = ttk.Treeview(tree_container, show="headings", selectmode="browse")
        vsb = ttk.Scrollbar(tree_container, orient="vertical", command=popup_tree.yview)
        hsb = ttk.Scrollbar(tree_container, orient="horizontal", command=popup_tree.xview)
        popup_tree.configure(yscrollcommand=vsb.set, xscrollcommand=hsb.set)

        popup_tree.grid(row=0, column=0, sticky="nsew")
        vsb.grid(row=0, column=1, sticky="ns")
        hsb.grid(row=1, column=0, sticky="ew")

        tree_container.grid_rowconfigure(0, weight=1)
        tree_container.grid_columnconfigure(0, weight=1)

        def sync_popup_tree(df):
            popup_tree.delete(*popup_tree.get_children())
            popup_tree["columns"] = list(df.columns)
            for col in df.columns:
                popup_tree.heading(col, text=col)
                popup_tree.column(col, width=130, minwidth=100, anchor=tk.W, stretch=False)
            for idx, row in df.iterrows():
                popup_tree.insert("", tk.END, iid=idx, values=list(row))

        sync_popup_tree(display_df)

        # Handle Task Execution within Analysis Pop-Up
        def execute_analysis_task(event):
            task_choice = analysis_task_var.get()
            if task_choice == "Choose Task...":
                return

            selected = popup_tree.selection()
            if not selected:
                messagebox.showwarning("Selection Required", "Please select a row in the Analysis table.", parent=popup)
                analysis_task_var.set("Choose Task...")
                return

            row_idx = int(selected[0])

            if task_choice == "Add Row As is it":
                self._task_add_row_as_is_it(row_idx)

            analysis_task_var.set("Choose Task...")

        analysis_task_cb.bind("<<ComboboxSelected>>", execute_analysis_task)

    def _task_add_row_as_is_it(self, row_idx):
        """
        Parses 'Source Unique IDs' from the selected formatted row, extracts 
        Document ID + Full Name components, looks up matching rows in the Original sheet,
        and exports them into a brand new sheet named 'As it is Data'.
        """
        if "Source Unique IDs" not in self.df_fmt.columns:
            messagebox.showerror("Missing Column", "The 'Source Unique IDs' column is missing from the Formatted sheet.")
            return

        unique_ids_val = str(self.df_fmt.at[row_idx, "Source Unique IDs"])
        if not unique_ids_val.strip():
            messagebox.showwarning("Empty Value", "The selected row has no 'Source Unique IDs' value to parse.")
            return

        # Split pipe-separated unique IDs (e.g. "GOL0000100475.xlsx_Hernandez, Adriana | ...")
        raw_parts = [p.strip() for p in unique_ids_val.split("|") if p.strip()]
        
        # Identify possible document/file ID and name columns in the original sheet
        orig_cols = list(self.df_orig.columns)
        doc_id_col = next((c for c in orig_cols if "document" in c.lower() or "sheet" in c.lower() or "id" in c.lower()), orig_cols[0] if orig_cols else None)
        name_col = next((c for c in orig_cols if "name" in c.lower()), orig_cols[1] if len(orig_cols) > 1 else None)

        if not doc_id_col or not name_col:
            messagebox.showerror("Matching Error", "Could not automatically identify Document ID or Name columns in the Original sheet.")
            return

        matched_rows = []
        for part in raw_parts:
            # Expected format in Unique ID: DocumentID_FullName or similar
            if "_" in part:
                doc_part, name_part = part.split("_", 1)
            else:
                doc_part, name_part = "", part

            # Search in Original Sheet
            match = self.df_orig[
                self.df_orig[doc_id_col].astype(str).str.contains(doc_part.strip(), case=False, regex=False) &
                self.df_orig[name_col].astype(str).str.contains(name_part.strip(), case=False, regex=False)
            ]
            if not match.empty:
                matched_rows.append(match)

        if not matched_rows:
            messagebox.showinfo("No Match Found", "No corresponding raw rows were found in the Original data sheet for these Unique IDs.")
            return

        target_df = pd.concat(matched_rows).drop_duplicates()

        # Save into Workbook as new sheet 'As it is Data' using openpyxl
        try:
            wb = openpyxl.load_workbook(self.file_path)
            sheet_title = "As it is Data"
            if sheet_title in wb.sheetnames:
                del wb[sheet_title]
            
            ws = wb.create_sheet(title=sheet_title)

            # Write Headers
            for c_idx, col_name in enumerate(target_df.columns):
                ws.cell(row=1, column=c_idx + 1, value=col_name)

            # Write Rows
            for r_idx, (_, row) in enumerate(target_df.iterrows()):
                for c_idx, val in enumerate(row):
                    ws.cell(row=r_idx + 2, column=c_idx + 1, value=val)

            wb.save(self.file_path)
            self.workbook_sheets = wb.sheetnames
            messagebox.showinfo("Success", f"Created new sheet '{sheet_title}' with {len(target_df)} raw unmerged row(s) successfully saved to workbook.")
            self.record_task("Add Row As is it", f"Exported {len(target_df)} original unmerged rows into new sheet 'As it is Data'")
        except Exception as e:
            messagebox.showerror("Export Error", f"Failed to create 'As it is Data' sheet in workbook:\n{e}")

    def _task_batch_pipe_extract(self, keep_index=0):
        popup = tk.Toplevel(self)
        title = "Keep First Pipe Value" if keep_index == 0 else "Keep Second Pipe Value"
        popup.title(title)
        popup.geometry("380x180")
        popup.grab_set()

        ttk.Label(popup, text=f"Select Column for {title}:").pack(pady=10)
        col_cb = ttk.Combobox(popup, values=list(self.df_fmt.columns), state="readonly")
        col_cb.set(self.df_fmt.columns[0])
        col_cb.pack(pady=5)

        def execute_extract():
            col = col_cb.get()
            undo_df = self.df_fmt.copy()
            undo_hl = copy.deepcopy(self.cell_highlights)

            changed = 0
            for r_idx in range(len(self.df_fmt)):
                orig_val = str(self.df_fmt.at[r_idx, col])
                if "|" in orig_val:
                    parts = [p.strip() for p in orig_val.split("|")]
                    if len(parts) > keep_index:
                        new_val = parts[keep_index]
                    else:
                        new_val = parts[0]

                    if orig_val != new_val:
                        self.df_fmt.at[r_idx, col] = new_val
                        self.cell_highlights[(r_idx, col)] = self.AUTO_YELLOW
                        changed += 1

            self.filtered_df_fmt = self.df_fmt.copy()
            self._sync_treeview(self.tree_fmt, self.filtered_df_fmt)
            self.record_task(title, f"Extracted part {keep_index+1} from Column '{col}' ({changed} cells updated & highlighted Yellow)", undo_df, undo_hl)
            messagebox.showinfo("Success", f"{title} executed on Column '{col}' ({changed} cells updated).")
            popup.destroy()

        ttk.Button(popup, text="Execute Batch Extract", command=execute_extract).pack(pady=15)

    def _task_fill_blanks(self):
        popup = tk.Toplevel(self)
        popup.title("Fill Blank Cells")
        popup.geometry("380x210")
        popup.grab_set()

        ttk.Label(popup, text="Target Column:").pack(pady=(10, 2))
        col_cb = ttk.Combobox(popup, values=list(self.df_fmt.columns), state="readonly")
        col_cb.set(self.df_fmt.columns[0])
        col_cb.pack(pady=2)

        ttk.Label(popup, text="Fill Empty Cells With Text:").pack(pady=(8, 2))
        ent = ttk.Entry(popup, width=30)
        ent.pack(pady=2)

        def execute_fill():
            col = col_cb.get()
            fill_val = ent.get()
            if not fill_val:
                return

            undo_df = self.df_fmt.copy()
            undo_hl = copy.deepcopy(self.cell_highlights)

            changed = 0
            for r_idx in range(len(self.df_fmt)):
                curr_val = str(self.df_fmt.at[r_idx, col]).strip()
                if curr_val == "":
                    self.df_fmt.at[r_idx, col] = fill_val
                    self.cell_highlights[(r_idx, col)] = self.AUTO_YELLOW
                    changed += 1

            self.filtered_df_fmt = self.df_fmt.copy()
            self._sync_treeview(self.tree_fmt, self.filtered_df_fmt)
            self.record_task("Fill Blank Cells", f"Filled blanks in Column '{col}' with '{fill_val}' ({changed} cells updated)", undo_df, undo_hl)
            messagebox.showinfo("Success", f"Filled {changed} empty cell(s) in Column '{col}'.")
            popup.destroy()

        ttk.Button(popup, text="Fill Blanks", command=execute_fill).pack(pady=15)

    def _task_remove_duplicates(self):
        undo_df = self.df_fmt.copy()
        undo_hl = copy.deepcopy(self.cell_highlights)

        before_count = len(self.df_fmt)
        self.df_fmt.drop_duplicates(inplace=True)
        self.df_fmt.reset_index(drop=True, inplace=True)
        after_count = len(self.df_fmt)
        removed_count = before_count - after_count

        self.filtered_df_fmt = self.df_fmt.copy()
        self._sync_treeview(self.tree_fmt, self.filtered_df_fmt)
        self.record_task("Remove Duplicates", f"Removed {removed_count} exact duplicate rows", undo_df, undo_hl)
        messagebox.showinfo("Success", f"Removed {removed_count} duplicate row(s) successfully.")

    def _task_trim_spaces(self):
        popup = tk.Toplevel(self)
        popup.title("Trim Extra Spaces")
        popup.geometry("380x180")
        popup.grab_set()

        ttk.Label(popup, text="Target Column for Space Trimming:").pack(pady=10)
        col_cb = ttk.Combobox(popup, values=list(self.df_fmt.columns), state="readonly")
        col_cb.set(self.df_fmt.columns[0])
        col_cb.pack(pady=5)

        def execute_trim():
            col = col_cb.get()
            undo_df = self.df_fmt.copy()
            undo_hl = copy.deepcopy(self.cell_highlights)

            changed = 0
            for r_idx in range(len(self.df_fmt)):
                orig_val = str(self.df_fmt.at[r_idx, col])
                trimmed = re.sub(r'\s+', ' ', orig_val.strip())
                if orig_val != trimmed:
                    self.df_fmt.at[r_idx, col] = trimmed
                    self.cell_highlights[(r_idx, col)] = self.AUTO_YELLOW
                    changed += 1

            self.filtered_df_fmt = self.df_fmt.copy()
            self._sync_treeview(self.tree_fmt, self.filtered_df_fmt)
            self.record_task("Trim Spaces", f"Trimmed spaces in Column '{col}' ({changed} cells updated & highlighted Yellow)", undo_df, undo_hl)
            messagebox.showinfo("Success", f"Trimmed spaces across Column '{col}' ({changed} cells changed).")
            popup.destroy()

        ttk.Button(popup, text="Trim Column Spaces", command=execute_trim).pack(pady=15)

    def _task_change_case(self):
        popup = tk.Toplevel(self)
        popup.title("Change Text Case")
        popup.geometry("380x220")
        popup.grab_set()

        ttk.Label(popup, text="Target Column:").pack(pady=(10, 2))
        col_cb = ttk.Combobox(popup, values=list(self.df_fmt.columns), state="readonly")
        col_cb.set(self.df_fmt.columns[0])
        col_cb.pack(pady=2)

        ttk.Label(popup, text="Select Case Transformation:").pack(pady=(10, 2))
        case_var = tk.StringVar(value="Proper Case")
        case_cb = ttk.Combobox(popup, textvariable=case_var, values=["UPPERCASE", "lowercase", "Proper Case"], state="readonly")
        case_cb.pack(pady=2)

        def execute_case_change():
            col = col_cb.get()
            case_type = case_var.get()
            undo_df = self.df_fmt.copy()
            undo_hl = copy.deepcopy(self.cell_highlights)

            changed = 0
            for r_idx in range(len(self.df_fmt)):
                orig_val = str(self.df_fmt.at[r_idx, col])
                if case_type == "UPPERCASE":
                    new_val = orig_val.upper()
                elif case_type == "lowercase":
                    new_val = orig_val.lower()
                elif case_type == "Proper Case":
                    new_val = orig_val.title()

                if orig_val != new_val:
                    self.df_fmt.at[r_idx, col] = new_val
                    self.cell_highlights[(r_idx, col)] = self.AUTO_YELLOW
                    changed += 1

            self.filtered_df_fmt = self.df_fmt.copy()
            self._sync_treeview(self.tree_fmt, self.filtered_df_fmt)
            self.record_task("Change Case", f"Converted Column '{col}' to {case_type} ({changed} cells highlighted Yellow)", undo_df, undo_hl)
            messagebox.showinfo("Success", f"Converted Column '{col}' to {case_type} ({changed} cells updated).")
            popup.destroy()

        ttk.Button(popup, text="Apply Case Conversion", command=execute_case_change).pack(pady=15)

    def _task_count_if(self):
        popup = tk.Toplevel(self)
        popup.title("CountIf Summary")
        popup.geometry("380x210")
        popup.grab_set()

        ttk.Label(popup, text="Target Column:").pack(pady=(10, 2))
        col_cb = ttk.Combobox(popup, values=list(self.df_fmt.columns), state="readonly")
        col_cb.set(self.df_fmt.columns[0])
        col_cb.pack(pady=2)

        ttk.Label(popup, text="Count values matching string:").pack(pady=(8, 2))
        ent = ttk.Entry(popup, width=30)
        ent.pack(pady=2)

        def execute_count_if():
            col = col_cb.get()
            query = ent.get().strip()
            if not query:
                return

            count = (self.df_fmt[col].astype(str).str.lower() == query.lower()).sum()
            msg = f"CountIf Summary for Column '{col}' where value == '{query}': {count} occurrences found."
            messagebox.showinfo("CountIf Result", msg)
            self.record_task("CountIf Summary", f"Counted '{query}' in Column '{col}' ({count} found)")

        ttk.Button(popup, text="Run CountIf Query", command=execute_count_if).pack(pady=15)

    def _task_find_data(self):
        popup = tk.Toplevel(self)
        popup.title("Find Data")
        popup.geometry("380x150")
        popup.grab_set()

        ttk.Label(popup, text="Search string in Formatted Data Sheet:").pack(pady=10)
        ent = ttk.Entry(popup, width=32)
        ent.pack(pady=4)

        def search():
            query = ent.get().strip()
            if not query:
                return
            results = []
            for r_idx, row in self.df_fmt.iterrows():
                for col in self.df_fmt.columns:
                    if query.lower() in str(row[col]).lower():
                        results.append(f"Row {r_idx + 1}, Column '{col}'")

            if results:
                summary = f"Found '{query}' in {len(results)} cell(s):\n" + "\n".join(results[:8])
                if len(results) > 8:
                    summary += f"\n...and {len(results) - 8} more."
                messagebox.showinfo("Find Results", summary)
                self.record_task("Find Data", f"Searched '{query}' ({len(results)} matches)")
            else:
                messagebox.showinfo("Find Results", f"No matches found for '{query}'.")

        ttk.Button(popup, text="Search Data", command=search).pack(pady=10)

    def _task_find_replace(self):
        popup = tk.Toplevel(self)
        popup.title("Find & Replace")
        popup.geometry("400x210")
        popup.grab_set()

        ttk.Label(popup, text="Find Text:").pack(anchor=tk.W, padx=25, pady=(10, 0))
        find_ent = ttk.Entry(popup)
        find_ent.pack(fill=tk.X, padx=25, pady=2)

        ttk.Label(popup, text="Replace With:").pack(anchor=tk.W, padx=25, pady=(6, 0))
        repl_ent = ttk.Entry(popup)
        repl_ent.pack(fill=tk.X, padx=25, pady=2)

        def replace():
            target = find_ent.get()
            replacement = repl_ent.get()
            if not target:
                return

            undo_df = self.df_fmt.copy()
            undo_hl = copy.deepcopy(self.cell_highlights)

            matches_count = 0
            for r_idx in range(len(self.df_fmt)):
                for col in self.df_fmt.columns:
                    orig_val = str(self.df_fmt.at[r_idx, col])
                    if target in orig_val:
                        new_val = orig_val.replace(target, replacement)
                        self.df_fmt.at[r_idx, col] = new_val
                        self.cell_highlights[(r_idx, col)] = self.AUTO_YELLOW
                        matches_count += 1

            self.filtered_df_fmt = self.df_fmt.copy()
            self._sync_treeview(self.tree_fmt, self.filtered_df_fmt)
            self.record_task("Find & Replace", f"Replaced '{target}' -> '{replacement}' ({matches_count} cells highlighted Yellow)", undo_df, undo_hl)
            messagebox.showinfo("Replace Complete", f"Replaced {matches_count} occurrence(s).")
            popup.destroy()

        ttk.Button(popup, text="Replace All", command=replace).pack(pady=15)

    def _task_insert_blank_row(self):
        row_idx = self._get_target_row()
        if row_idx is None:
            row_idx = len(self.df_fmt)

        undo_df = self.df_fmt.copy()
        undo_hl = copy.deepcopy(self.cell_highlights)

        blank_row = {col: "" for col in self.df_fmt.columns}

        top = self.df_fmt.iloc[:row_idx]
        bottom = self.df_fmt.iloc[row_idx:]

        self.df_fmt = pd.concat([top, pd.DataFrame([blank_row]), bottom], ignore_index=True)
        for col in self.df_fmt.columns:
            self.cell_highlights[(row_idx, col)] = self.AUTO_YELLOW

        self.filtered_df_fmt = self.df_fmt.copy()
        self._sync_treeview(self.tree_fmt, self.filtered_df_fmt)
        self.record_task("Insert Blank Row", f"Inserted blank row at index {row_idx + 1} (Highlighted Yellow)", undo_df, undo_hl)

    def _task_copy_cell_value(self):
        popup = tk.Toplevel(self)
        popup.title("Copy Cell Value")
        popup.geometry("400x260")
        popup.grab_set()

        ttk.Label(popup, text="Source Cell:", font=("Helvetica", 9, "bold")).pack(pady=5)
        src_f = ttk.Frame(popup)
        src_f.pack(fill=tk.X, padx=20)
        ttk.Label(src_f, text="Row:").pack(side=tk.LEFT)
        src_r = ttk.Entry(src_f, width=6)
        src_r.insert(0, "1")
        src_r.pack(side=tk.LEFT, padx=5)
        ttk.Label(src_f, text="Column:").pack(side=tk.LEFT, padx=5)
        src_c = ttk.Combobox(src_f, values=list(self.df_fmt.columns), state="readonly", width=15)
        src_c.set(self.df_fmt.columns[0])
        src_c.pack(side=tk.LEFT)

        ttk.Label(popup, text="Target Cell:", font=("Helvetica", 9, "bold")).pack(pady=(15, 5))
        tgt_f = ttk.Frame(popup)
        tgt_f.pack(fill=tk.X, padx=20)
        ttk.Label(tgt_f, text="Row:").pack(side=tk.LEFT)
        tgt_r = ttk.Entry(tgt_f, width=6)
        tgt_r.insert(0, "1")
        tgt_r.pack(side=tk.LEFT, padx=5)
        ttk.Label(tgt_f, text="Column:").pack(side=tk.LEFT, padx=5)
        tgt_c = ttk.Combobox(tgt_f, values=list(self.df_fmt.columns), state="readonly", width=15)
        tgt_c.set(self.df_fmt.columns[0])
        tgt_c.pack(side=tk.LEFT)

        def copy_val():
            try:
                sr, tr = int(src_r.get()) - 1, int(tgt_r.get()) - 1
                sc, tc = src_c.get(), tgt_c.get()

                undo_df = self.df_fmt.copy()
                undo_hl = copy.deepcopy(self.cell_highlights)

                val = self.df_fmt.at[sr, sc]
                self.df_fmt.at[tr, tc] = val
                self.cell_highlights[(tr, tc)] = self.AUTO_YELLOW

                self.filtered_df_fmt = self.df_fmt.copy()
                self._sync_treeview(self.tree_fmt, self.filtered_df_fmt)
                self.record_task("Copy Cell Value", f"Copied from (R{sr+1}, '{sc}') to (R{tr+1}, '{tc}') (Highlighted Yellow)", undo_df, undo_hl)
                popup.destroy()
            except Exception as e:
                messagebox.showerror("Execution Error", f"Invalid cell coordinates: {e}")

        ttk.Button(popup, text="Execute Copy", command=copy_val).pack(pady=18)

    def _task_view_history(self):
        popup = tk.Toplevel(self)
        popup.title("View Task History")
        popup.geometry("550x340")
        popup.grab_set()

        ttk.Label(popup, text="Recorded Operations History Log", font=("Helvetica", 11, "bold")).pack(pady=10)

        lb = tk.Listbox(popup, font=("Consolas", 9))
        lb.pack(fill=tk.BOTH, expand=True, padx=15, pady=10)

        for item in self.history:
            lb.insert(tk.END, f"[{item['task']}] - {item['details']}")

    # =========================================================================
    # 8. STYLE-PRESERVING EXPORT PIPELINE
    # =========================================================================
    def export_final_file(self):
        save_path = filedialog.asksaveasfilename(defaultextension=".xlsx", filetypes=[("Excel Workbooks", "*.xlsx")])
        if not save_path:
            return

        try:
            wb = openpyxl.load_workbook(self.file_path)

            if self.fmt_sheet_name in wb.sheetnames:
                ws = wb[self.fmt_sheet_name]

                # Update modified cell values
                for r_idx, row in self.df_fmt.iterrows():
                    for c_idx, col_name in enumerate(self.df_fmt.columns):
                        val = row[col_name]
                        ws.cell(row=r_idx + 2, column=c_idx + 1, value=val)

                # Apply recorded highlights (including automatic yellow for modified cells)
                for (r_idx, col_name), hex_color in self.cell_highlights.items():
                    if col_name in list(self.df_fmt.columns):
                        c_idx = list(self.df_fmt.columns).index(col_name) + 1
                        clean_hex = hex_color.lstrip("#")
                        if len(clean_hex) == 6:
                            fill = PatternFill(start_color=clean_hex, end_color=clean_hex, fill_type="solid")
                            ws.cell(row=r_idx + 2, column=c_idx).fill = fill

            wb.save(save_path)
            messagebox.showinfo("Export Successful", f"Workbook exported successfully to:\n{save_path}\n\nAll modified cells are highlighted in Yellow.")
        except Exception as e:
            messagebox.showerror("Export Failed", f"An error occurred while exporting the workbook:\n{e}")

if __name__ == "__main__":
    app = AnalysisApp()
    app.mainloop()
